"""Sensitive Data Scanner — evolução do `pii_detection` para escanear
documentos inteiros (`.txt`, `.pdf`, `.docx`), não só texto solto.

Não reimplementa NENHUMA lógica de detecção de PII: extrai o texto do
documento (a única responsabilidade nova deste módulo) e repassa para
`pii_detection.detect()`, que já é o motor real e testado do V1. Isso segue a
mesma regra de composição documentada pelo `ripd_engine`.

Formatos suportados nesta versão: `.txt` (leitura direta), `.pdf` (via
`pypdf`) e `.docx` (via `python-docx`). OCR de imagem/PDF escaneado (item
citado no ROADMAP original como "documentos/OCR") fica como TODO explícito —
ver `CHANGELOG.md` deste módulo: exige um motor de OCR (ex. Tesseract) que
não está nas dependências do projeto nesta versão.
"""
from __future__ import annotations

from pathlib import Path

from core.pii_detection.detector import detect
from shared.schemas import DocumentScanResult

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


def _extract_txt(path: Path) -> tuple[str, int | None, str]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return text, None, "Texto lido diretamente (UTF-8, erros substituídos por replacement char)."


def _extract_pdf(path: Path) -> tuple[str, int | None, str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages_text = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages_text)
    notes = f"Texto extraído via pypdf de {len(reader.pages)} página(s)."
    if not text.strip():
        notes += " Nenhum texto extraível encontrado — provável PDF escaneado (imagem), requer OCR (não suportado nesta versão)."
    return text, len(reader.pages), notes


def _extract_docx(path: Path) -> tuple[str, int | None, str]:
    import docx

    document = docx.Document(str(path))
    paragraphs = [p.text for p in document.paragraphs]
    tables_text = [
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    ]
    text = "\n".join(paragraphs + tables_text)
    notes = f"Texto extraído via python-docx: {len(paragraphs)} parágrafo(s) + {len(document.tables)} tabela(s)."
    return text, None, notes


_EXTRACTORS = {
    ".txt": _extract_txt,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
}


def scan_document(file_path: str | Path) -> DocumentScanResult:
    """Extrai o texto de um documento e escaneia por PII/dado sensível.

    Args:
        file_path: caminho para um arquivo `.txt`, `.pdf` ou `.docx`.

    Returns:
        DocumentScanResult com o `PIIDetectionResult` real (via
        `pii_detection.detect`) sobre o texto extraído.

    Levanta:
        FileNotFoundError: se `file_path` não existir.
        ValueError: se a extensão não for suportada (ver `SUPPORTED_EXTENSIONS`).
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Documento não encontrado: {path}")

    extension = path.suffix.lower()
    if extension not in _EXTRACTORS:
        raise ValueError(
            f"Extensão '{extension}' não suportada. Suportadas: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    text, pages, notes = _EXTRACTORS[extension](path)
    pii_result = detect(text)

    return DocumentScanResult(
        file_name=path.name,
        file_type=extension.lstrip("."),
        pages_scanned=pages,
        characters_extracted=len(text),
        pii_result=pii_result,
        extraction_notes=notes,
    )
